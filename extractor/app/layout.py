"""AI layout reading: the AI labels rows, the code reads the numbers.

Used when the rule-based readers can't make sense of a BOQ's layout. The
file becomes a table of rows — one per Excel row or PDF text line — with
cells keyed by column. The web app sends a masked copy of the table to the
AI provider, in which every standalone number is replaced by "#", and gets
back:

- which column holds the reference, description, quantity, unit, rate and
  amount (asked only when the column headers don't say), and
- one label per row (LABELS below).

The BOQ is then rebuilt here from the original file, so every quantity,
rate and amount comes from the BOQ itself: the AI never sees them and can't
change them. The web app passes back a hash of the masked table it labelled,
so labels can never be applied to a different reading of the file.
"""

import hashlib
import json
import re
from dataclasses import dataclass, field

import pdfplumber
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field

from .excel_reader import _num, _text, find_header, inspect_excel, open_workbook, visible_sheets
from .hierarchy import BillStart, Event, ItemRow, TextBlock, TotalRow
from .models import Cover, Issue, PageTotal
from .numbers import parse_number, rate_note
from .pdf_reader import _BILL, SAME_LINE, ExtractError, Line, Word, _group_lines, _words, inspect_pdf

LABELS = {
    "B": "bill title",
    "H": "heading",
    "M": "main description",
    "I": "item",
    "C": "item continuation",
    "T": "page total",
    "X": "ignore",
}
ROLES = ("reference", "description", "quantity", "unit", "rate", "amount")
TOP = "top"  # PDF text above the column headers: page and bill titles
MAX_ROWS = 15000
MAX_LAYOUTS = 30
_SAMPLE_TEXT_ROWS, _SAMPLE_NUMBER_ROWS = 10, 30


@dataclass
class LayoutRow:
    id: int
    page: int
    layout: str
    bold: bool
    cells: dict[str, str]  # column key -> text exactly as in the file
    bill_hint: str | None = None  # Excel: the sheet name, which starts a new bill


@dataclass
class Layout:
    key: str
    headers: dict[str, str]  # column key -> header label
    roles: dict[str, str | None] | None  # role -> column key; None = ask the AI


@dataclass
class RowTable:
    file_type: str
    pages: int
    layouts: dict[str, Layout] = field(default_factory=dict)
    rows: list[LayoutRow] = field(default_factory=list)

    def add(self, row: LayoutRow) -> None:
        if len(self.rows) >= MAX_ROWS:
            raise ExtractError(
                "TOO_LARGE_FOR_AI",
                f"This BOQ has more than {MAX_ROWS:,} rows, which is too many to read with AI help. Split it into smaller files.",
            )
        self.rows.append(row)


# ── Masking ───────────────────────────────────────────────────────────────

_EDGE = "()[]{},;:$£€"


def _is_number(token: str) -> bool:
    core = token.strip(_EDGE).rstrip(".")
    return bool(core) and parse_number(core) is not None


def mask(text: str) -> str:
    """Every standalone number becomes "#". Numbers inside words ("500mm",
    "P4") and codes ("2.5.1") stay, so the AI can still read the text."""
    return re.sub(r"\S+", lambda m: "#" if _is_number(m.group()) else m.group(), text)


def masked_cells(row: LayoutRow) -> dict[str, str]:
    return {k: mask(v) for k, v in row.cells.items()}


def rows_hash(table: RowTable) -> str:
    data = [[r.id, r.page, r.layout, r.bold, masked_cells(r)] for r in table.rows]
    return hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


# ── Column roles from header labels ───────────────────────────────────────


def guess_roles(headers: dict[str, str]) -> dict[str, str | None] | None:
    """Roles from header labels, or None when they don't clearly say."""
    up = {k: v.upper().strip(" .:") for k, v in headers.items()}

    def first(pattern: str, exclude: str | None = None) -> str | None:
        return next((k for k, v in up.items() if re.search(pattern, v) and not (exclude and exclude in v)), None)

    roles = {
        # Labels printed close together can merge ("ITEMNO").
        "reference": first(r"^(ITEM|ITEM ?NO|ITEM ?REF|REF|REF ?NO|NO|S\.? ?NO|SN|SR|CODE)\.?$"),
        "description": first(r"DESCRIPTION|PARTICULARS"),
        "quantity": first(r"\b(QTY|QUANTITY|QUANTITIES)\b"),
        "unit": first(r"\b(UNIT|UNITS|UOM)\b", exclude="RATE"),
        "rate": first(r"RATE"),
        "amount": first(r"\b(AMOUNT|TOTAL)\b"),
    }
    chosen = [v for v in roles.values() if v]
    if not roles["description"] or not (roles["quantity"] or roles["amount"]) or len(set(chosen)) != len(chosen):
        return None
    return roles


# ── PDF ───────────────────────────────────────────────────────────────────

_PDF_HEADER_WORD = re.compile(r"^(QTY|QUANTITY|QUANTITIES|UNIT|UNITS|UOM|RATE|AMOUNT|TOTAL)$")
_HEADER_VOCAB = re.compile(r"^(QTY|QUANTITY|QUANTITIES|UNIT|UNITS|UOM|RATE|AMOUNT|TOTAL|NR|NO|REF|ITEM|DESCRIPTION|REMARKS|SUM|COST|PRICE|AED|SAR|QAR|USD)$")
_FRAGMENT = re.compile(r"^[\d,.()\-]+$")
HEADER_STACK = 14.0  # pt: header lines this close above / below belong to the same header


def _is_text_column(label: str) -> bool:
    """Left-aligned text columns, whose text starts left of their centred header."""
    return bool(re.search(r"DESCRIPTION|PARTICULARS|REMARKS|NOTES|COMMENTS", label.upper()))


def _header_like(line: Line) -> bool:
    return any(_HEADER_VOCAB.match(w.text.upper().strip(":.()")) for w in line.words)


@dataclass
class _Header:
    top: float  # first and last line of the header block
    bottom: float
    labels: list[tuple[str, float, float]]  # (text, x0, x1), left to right


def _pdf_header(lines: list[Line]) -> _Header | None:
    """The column header. Headers can be stacked over several lines
    ("AMOUNT" / "(1 NR.)"), so header-like lines just above and below the
    main one are included, and words that overlap horizontally are one label."""
    for i, line in enumerate(lines[:15]):
        texts = [w.text.upper().strip(":.()") for w in line.words]
        if not any(re.search(r"DESCRIPTION|PARTICULARS", t) for t in texts) or sum(1 for t in texts if _PDF_HEADER_WORD.match(t)) < 2:
            continue
        block = [line]
        for step in (-1, 1):
            j = i + step
            while 0 <= j < len(lines) and abs(lines[j].top - block[-1 if step > 0 else 0].top) <= HEADER_STACK and _header_like(lines[j]):
                block.insert(len(block) if step > 0 else 0, lines[j])
                j += step
        words = sorted((w for l in block for w in l.words), key=lambda w: w.x0)
        groups: list[list[Word]] = []
        for w in words:
            if groups and w.x0 - max(g.x1 for g in groups[-1]) < 6:  # overlapping or touching
                groups[-1].append(w)
            else:
                groups.append([w])
        labels = [(" ".join(g.text for g in sorted(grp, key=lambda g: (g.top, g.x0))), min(g.x0 for g in grp), max(g.x1 for g in grp)) for grp in groups]
        return _Header(top=block[0].top, bottom=block[-1].top, labels=labels)
    return None


def _bands(labels: list[tuple[str, float, float]]) -> list[tuple[str, float]]:
    """Column keys and their left edges. Numbers are right-aligned and can
    start left of their header, so a column begins a little before it; a
    description starts just after the column before it, because its text
    starts well left of its centred header."""
    out = [("c0", float("-inf"))]
    for i in range(1, len(labels)):
        prev, cur = labels[i - 1], labels[i]
        edge = prev[2] + 5 if _is_text_column(cur[0]) else cur[1] - min(15.0, (cur[1] - prev[2]) * 0.4)
        out.append((f"c{i}", edge))
    return out


def _cell_text(words: list[Word]) -> str:
    """Joins a cell's words. Some PDFs split one number into touching pieces
    ("2" + "00.00", "3" + ",500.00"); those are glued back together, and so
    is a superscript after a unit ("m" + "2")."""
    out, prev = "", None
    for w in words:
        touching = prev is not None and w.x0 - prev.x1 <= 1.5
        if touching and _FRAGMENT.match(prev.text) and _FRAGMENT.match(w.text):  # type: ignore[union-attr]
            out += w.text
        elif prev is not None and w.text.isdigit() and w.size < prev.size * 0.8 and w.x0 - prev.x1 <= 3:
            out += w.text
        else:
            out = f"{out} {w.text}" if out else w.text
        prev = w
    return out


def _pdf_table(path: str) -> RowTable:
    with pdfplumber.open(path) as pdf:
        if pdf.metadata.get("Encrypted"):
            raise ExtractError("ENCRYPTED", "This PDF is password-protected. Remove the password and upload it again.")
        table = RowTable(file_type="pdf", pages=len(pdf.pages))
        if sum(len(p.chars) for p in pdf.pages[: min(5, table.pages)]) < 50:
            raise ExtractError("NO_TEXT", "This PDF has no readable text — it looks like a scanned document. Upload the original PDF or Excel BOQ.")
        bands: list[tuple[str, float]] | None = None
        key = ""
        previous_top: set[str] = set()
        for page_no, page in enumerate(pdf.pages, start=1):
            lines = _group_lines(_words(page))
            header = _pdf_header(lines)
            if header:
                labels = header.labels
                bands = _bands(labels)
                key = "pdf:" + "|".join(text.upper() for text, _, _ in labels)
                if key not in table.layouts:
                    headers = {k: text for (k, _), (text, _, _) in zip(bands, labels)}
                    table.layouts[key] = Layout(key, headers, guess_roles(headers))
                top = [l for l in lines if l.top < header.top - SAME_LINE]
                body = [l for l in lines if l.top > header.bottom + SAME_LINE]
                # Page titles repeat on every page; only a change is news.
                for l in top:
                    if l.text not in previous_top:
                        table.add(LayoutRow(len(table.rows), page_no, key, l.bold, {TOP: l.text}))
                previous_top = {l.text for l in top}
            elif bands is None:
                continue  # cover / summary pages before the first table
            else:
                body = lines  # a continuation page without repeated headers
            for line in body:
                by_band: dict[str, list[Word]] = {}
                for w in sorted(line.words, key=lambda w: w.x0):
                    band = next(k for k, edge in reversed(bands) if w.x0 >= edge)
                    by_band.setdefault(band, []).append(w)
                cells = {k: _cell_text(by_band[k]) for k, _ in bands if k in by_band}
                table.add(LayoutRow(len(table.rows), page_no, key, line.bold, cells))
    if len(table.layouts) > MAX_LAYOUTS:
        raise ExtractError("TOO_MANY_LAYOUTS", "This PDF changes its column layout too often to read with AI help.")
    return table


# ── Excel ─────────────────────────────────────────────────────────────────


def _xlsx_table(path: str, mapping: dict | None) -> RowTable:
    sheets = visible_sheets(open_workbook(path))
    table = RowTable(file_type="xlsx", pages=len(sheets))
    for sheet_no, ws in enumerate(sheets, start=1):
        cols = find_header(ws, mapping=mapping)
        if cols:
            headers = {get_column_letter(c.column): _text(c.value) for c in ws[cols.header_row] if _text(c.value)}
            letter = lambda col: get_column_letter(col) if col else None  # noqa: E731
            roles = {"reference": letter(cols.ref), "description": letter(cols.desc), "quantity": letter(cols.qty), "unit": letter(cols.unit), "rate": letter(cols.rate), "amount": letter(cols.amount)}
            key = "xlsx:" + "|".join(f"{k}={v}" for k, v in headers.items())
        else:
            headers, roles, key = {}, None, f"sheet:{ws.title.strip()}"
        table.layouts.setdefault(key, Layout(key, headers, roles))
        for row in ws.iter_rows():
            if cols and row[0].row == cols.header_row:
                continue
            cells: dict[str, str] = {}
            bold = False
            for c in row:
                v = c.value
                if v is None:
                    continue
                t = _num(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else _text(v)
                if t:
                    cells[get_column_letter(c.column)] = t
                    bold = bold or bool(getattr(c, "font", None) is not None and c.font.b)
            if cells:
                table.add(LayoutRow(len(table.rows), sheet_no, key, bold, cells, bill_hint=ws.title.strip()))
    if len(table.layouts) > MAX_LAYOUTS:
        raise ExtractError("TOO_MANY_LAYOUTS", "This workbook has too many differently laid-out sheets to read with AI help.")
    return table


def build_table(path: str, file_type: str, mapping: dict | None = None) -> RowTable:
    return _pdf_table(path) if file_type == "pdf" else _xlsx_table(path, mapping)


def _columns(table: RowTable, key: str) -> set[str]:
    """Column keys a role may point at: named in the header or holding data."""
    return set(table.layouts[key].headers) | {k for r in table.rows if r.layout == key for k in r.cells if k != TOP}


def describe(table: RowTable) -> dict:
    """The masked table for the web app (POST /v1/rows). No raw numbers."""
    masked = [masked_cells(r) for r in table.rows]
    layouts = []
    for key, layout in table.layouts.items():
        mine = [r for r in table.rows if r.layout == key]
        text_rows = [r.id for r in mine[:_SAMPLE_TEXT_ROWS]]
        number_rows = [r.id for r in mine if "#" in masked[r.id].values() and r.id not in text_rows][:_SAMPLE_NUMBER_ROWS]
        layouts.append(
            {
                "key": key,
                "headers": layout.headers,
                "roles": layout.roles,
                "columns": sorted(_columns(table, key)),
                "sample": [masked[i] for i in sorted(text_rows + number_rows)],
            }
        )
    return {
        "file_type": table.file_type,
        "pages": table.pages,
        "rows_hash": rows_hash(table),
        "layouts": layouts,
        "rows": [{"id": r.id, "page": r.page, "layout": r.layout, "bold": r.bold, "cells": masked[r.id]} for r in table.rows],
    }


# ── Rebuilding the BOQ from the labels ────────────────────────────────────


class LayoutPlan(BaseModel):
    """What the web app sends back: column roles per layout and one label per row."""

    model_config = ConfigDict(extra="forbid")
    rows_hash: str = Field(min_length=64, max_length=64)
    roles: dict[str, dict[str, str | None]]
    labels: str = Field(max_length=MAX_ROWS)


def _check_plan(table: RowTable, plan: LayoutPlan) -> None:
    def bad(message: str):
        return ExtractError("INVALID_LAYOUT", message)

    if plan.rows_hash != rows_hash(table):
        raise bad("The file was read differently from when the AI looked at it. Run the extraction again.")
    if len(plan.labels) != len(table.rows) or set(plan.labels) - LABELS.keys():
        raise bad("The row labels don't match this file.")
    for key in {r.layout for r in table.rows}:
        roles = plan.roles.get(key)
        columns = _columns(table, key)
        if roles is None or set(roles) - set(ROLES) or not roles.get("description"):
            raise bad("A column layout is missing its description column.")
        chosen = [v for v in roles.values() if v]
        if len(set(chosen)) != len(chosen) or set(chosen) - columns:
            raise bad("The column roles don't match this file.")


def _has_letters(text: str | None) -> bool:
    return bool(text and re.search(r"[A-Za-z]{2,}", text))


def _leading_value(text: str) -> str:
    """A number cell sometimes catches the start of the next column's text
    ("3,500.00 Assumed wood"); keep just the number or note at its start."""
    if not text or parse_number(text) is not None or rate_note(text):
        return text
    first = text.split()[0]
    return first if parse_number(first) is not None or rate_note(first) else text


@dataclass
class LayoutRead:
    events: list[Event]
    cover: Cover
    page_totals: list[PageTotal]
    issues: list[Issue]
    pages: int


def assemble(table: RowTable, plan: LayoutPlan) -> tuple[list[Event], list[Issue]]:
    events: list[Event] = []
    issues: list[Issue] = []
    skipped: list[LayoutRow] = []
    bills = 0
    hint: str | None = None
    item: ItemRow | None = None
    main: list[str] = []
    main_page = 0

    def start_bill(title: str) -> None:
        nonlocal bills
        bills += 1
        m = _BILL.search(title)
        if events and isinstance(events[-1], BillStart):
            events.pop()  # a bill with nothing in it: replaced by this one
        events.append(BillStart(bill_no=m.group(1).zfill(2) if m else f"{bills:02d}", title=(m.group(2).strip() if m else title)[:300]))

    def flush_main() -> None:
        if main:
            events.append(TextBlock([(t, False) for t in main], main_page, kind="main"))
            main.clear()

    for row, label in zip(table.rows, plan.labels):
        roles = plan.roles[row.layout]

        def cell(role: str) -> str:
            key = roles.get(role)
            value = row.cells.get(key, "").strip() if key else ""
            return _leading_value(value) if role in ("quantity", "rate", "amount") else value

        desc, unit = cell("description"), cell("unit") or None
        if row.bill_hint != hint:
            flush_main()
            item, hint = None, row.bill_hint
            if hint:
                start_bill(hint)
        if label == "C" and item is None:
            label = "M"
        if label in "XBHM" and parse_number(cell("quantity")) is not None and unit and re.search(r"[A-Za-z]", unit):
            skipped.append(row)

        if label == "X":
            continue
        if label == "C":
            if desc:
                item.lines.append(desc)  # type: ignore[union-attr]
            for attr, role in (("qty_raw", "quantity"), ("rate_raw", "rate"), ("amount_raw", "amount")):
                if getattr(item, attr) is None and cell(role):
                    setattr(item, attr, cell(role))
            if item.unit_raw is None:  # type: ignore[union-attr]
                item.unit_raw = unit  # type: ignore[union-attr]
            if item.ref == "-" and cell("reference") and not _has_letters(cell("reference")):  # type: ignore[union-attr]
                item.ref = cell("reference")  # type: ignore[union-attr]
            continue

        item = None
        text = next((t for t in (desc, cell("reference"), row.cells.get(TOP, ""), *row.cells.values()) if _has_letters(t)), "")
        if label == "M":
            if not main:
                main_page = row.page
            if text:
                main.append(text)
            continue
        flush_main()
        if label == "B":
            start_bill(f"{hint} — {text}" if hint and text else text or "Bill of Quantities")
        elif label == "H":
            if text:
                events.append(TextBlock([(text, True)], row.page, kind="heading"))
        elif label == "I":
            ref = cell("reference")
            item = ItemRow(
                ref=ref if ref and ref != "0" else "-",
                lines=[desc] if desc else [],
                page=row.page,
                qty_raw=cell("quantity") or None,
                unit_raw=unit,
                rate_raw=cell("rate") or None,
                amount_raw=cell("amount") or None,
            )
            events.append(item)
        elif label == "T":
            amount = cell("amount")
            if parse_number(amount) is None:
                amount = next((v for v in reversed(row.cells.values()) if parse_number(v) is not None), "")
            events.append(TotalRow(amount_raw=amount or None, page=row.page))
    flush_main()

    for row in skipped[:20]:
        text = " ".join(row.cells.values())[:80]
        issues.append(Issue(code="PRICED_ROW_SKIPPED", severity="warning", message=f'Page {row.page}: a row with a quantity and unit wasn\'t read as an item ("{text}"). Check it.', page=row.page))
    if len(skipped) > 20:
        issues.append(Issue(code="PRICED_ROW_SKIPPED", severity="warning", message=f"{len(skipped) - 20} more rows with a quantity and unit weren't read as items.", page=None))
    return events, issues


def read_with_layout(path: str, file_type: str, mapping: dict | None, layout: dict) -> LayoutRead:
    table = build_table(path, file_type, mapping)
    plan = LayoutPlan.model_validate(layout)
    _check_plan(table, plan)
    events, issues = assemble(table, plan)
    issues.insert(0, Issue(code="AI_LAYOUT", severity="info", message="This BOQ's layout was read with AI help. Every number comes from the file, but check headings, main descriptions and which lines belong to which item.", page=None))
    cover = inspect_pdf(path)[0] if file_type == "pdf" else inspect_excel(path)[0]
    return LayoutRead(events=events, cover=cover, page_totals=[], issues=issues, pages=table.pages)
