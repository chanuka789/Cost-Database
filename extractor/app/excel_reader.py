"""Reads an Excel (.xlsx) BOQ into hierarchy events (see hierarchy.py).

Each visible sheet with a column header row (DESCRIPTION / QTY / UNIT …) is
read as one bill. Sheets without one (cover, summary) are skipped, except
that the first is used for cover guesses (project name, date, stage).

Rows:
- a row with an item reference or a quantity          -> an item
- description-only rows directly above an item whose own
  description cell is empty                           -> that item's text
- other description-only rows, grouped until a blank row -> a text block
  (heading or main description); bold cells read as bold
- "to collection" / page total rows, in any column     -> a total, checked
  against the items above it
- carried / brought forward, other total rows and
  repeated column-header rows                          -> ignored

The header row can sit well down the sheet (Aldar BOQs put a title block
of ~50 rows above it). A reference of 0 is a formula artefact, not a ref.
"""

import re
from dataclasses import dataclass, replace
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from .hierarchy import BillStart, Event, ItemRow, TextBlock, TotalRow
from .models import Cover, Issue, PageTotal
from .pdf_reader import ExtractError, guess_date, guess_stage

_BILL = re.compile(r"\bBILL\s*(?:NO\.?|NUMBER)?\s*[:.]?\s*(\d+[A-Z]?)\s*[-–:]\s*(.+)$", re.I)
_SKIP_ROW = re.compile(r"CARRIED\s+(TO|FORWARD)|BROUGHT\s+FORWARD|TO\s+COLLECTION|^\s*(SUB-?\s*)?TOTAL\b|^\s*COLLECTION\s*$", re.I)
# Page sums: the items since the previous one should add up to these.
_PAGE_TOTAL = re.compile(r"TO\s+COLLECTION|PAGE\s+TOTAL", re.I)
# Safe to look for in every cell, not just the description.
_SKIP_ANY_CELL = re.compile(r"CARRIED\s+(TO|FORWARD)|BROUGHT\s+FORWARD|TO\s+COLLECTION|PAGE\s+TOTAL", re.I)
HEADER_SEARCH_ROWS = 150
_WORD = re.compile(r"[A-Za-z]{4,}")  # item references are short codes, never words


@dataclass
class _Cols:
    ref: int | None
    desc: int
    qty: int
    unit: int
    rate: int | None
    amount: int | None
    header_row: int


@dataclass
class ExcelRead:
    events: list[Event]
    cover: Cover
    page_totals: list[PageTotal]
    issues: list[Issue]
    pages: int


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _num(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return str(Decimal(str(value)))
    return str(value).strip()


def find_header(ws, max_rows: int = HEADER_SEARCH_ROWS, mapping: dict | None = None) -> _Cols | None:
    """The column header row. Rows of working columns ("ADJ. QTY", "UNIT",
    "Changes in Qty / Description") can look like a header too, so every
    candidate is scored and the one that names the most BOQ columns wins."""
    best: tuple[int, _Cols] | None = None
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows), start=1):
        if mapping:
            exact = {}
            for cell in row:
                label = _text(cell.value)
                if label:
                    exact.setdefault(label, []).append(cell.column)
            if all(v is None or len(exact.get(v, [])) == 1 for v in mapping.values()):
                def mapped(field):
                    label = mapping[field]
                    return exact[label][0] if label else None
                return _Cols(ref=mapped("reference"), desc=mapped("description"), qty=mapped("quantity"), unit=mapped("unit"), rate=mapped("rate"), amount=mapped("amount"), header_row=r)
            continue
        labels = {c.column: _text(c.value).upper() for c in row if _text(c.value)}

        def col(*names: str) -> int | None:
            for k, v in labels.items():
                if any(v == n or v.startswith(n + " ") or v.endswith(" " + n) or v == n + "." for n in names):
                    return k
            return None

        desc = next((k for k, v in labels.items() if "DESCRIPTION" in v), None)
        qty = col("QTY", "QUANTITY", "QTY.")
        unit = col("UNIT", "UOM", "UNITS")
        if desc and qty and unit:
            ref = col("ITEM", "ITEM NO", "REF", "NO", "S.NO", "SN", "ITEM REF")
            if ref == desc:
                ref = None
            found = _Cols(ref=ref, desc=desc, qty=qty, unit=unit, rate=col("RATE", "UNIT RATE"), amount=col("AMOUNT", "TOTAL", "AMOUNT (SAR)", "AMOUNT (AED)", "AMOUNT (QAR)"), header_row=r)
            score = sum(1 for c in (found.ref, found.rate, found.amount) if c) + (labels[desc].startswith("DESCRIPTION") or labels[desc] == "ITEM DESCRIPTION")
            if best is None or score > best[0]:
                best = (score, found)
    return best[1] if best else None


def _is_item_row(values: dict[int, object], cols: _Cols) -> bool:
    desc, qty, unit = values.get(cols.desc), values.get(cols.qty), values.get(cols.unit)
    return (
        isinstance(desc, str) and len(desc.strip()) >= 3
        and isinstance(qty, (int, float)) and not isinstance(qty, bool)
        and isinstance(unit, str) and 0 < len(unit.strip()) <= 12
    )


def _fit_count(ws, cols: _Cols, limit: int = 200) -> int:
    """How many rows have text, a number and a unit where `cols` expects them (up to `limit`)."""
    found = 0
    for row in ws.iter_rows(min_row=cols.header_row + 1, values_only=False):
        if _is_item_row({c.column: c.value for c in row}, cols):
            found += 1
            if found >= limit:
                break
    return found


def _item_like_rows(ws) -> int:
    """Rows with some text and at least two numbers — a rough sign of BOQ items."""
    count = 0
    for row in ws.iter_rows(values_only=True):
        numbers = sum(1 for v in row if isinstance(v, (int, float)) and not isinstance(v, bool))
        if numbers >= 2 and any(isinstance(v, str) and len(v.strip()) >= 3 for v in row):
            count += 1
    return count


def open_workbook(path: str):
    try:
        return load_workbook(path, data_only=True)
    except InvalidFileException as e:
        raise ExtractError("UNSUPPORTED", "This Excel file can't be opened. Save it as .xlsx (Excel Workbook) and upload again.") from e
    except Exception as e:  # corrupt or password-protected
        raise ExtractError("UNREADABLE", "This Excel file can't be read. If it's password-protected, remove the password and upload again.") from e


def visible_sheets(wb) -> list:
    return [ws for ws in wb.worksheets if ws.sheet_state == "visible"]


def _cover_from(ws) -> Cover:
    text = " ".join(_text(c.value) for row in ws.iter_rows(max_row=60) for c in row if _text(c.value))
    first = next((_text(c.value) for row in ws.iter_rows(max_row=60) for c in row if _text(c.value)), None)
    stage_text, stage_guess = guess_stage(text)
    return Cover(project_name=first, boq_date=guess_date(text), stage_text=stage_text, stage_guess=stage_guess)


def inspect_excel(path: str) -> tuple[Cover, int]:
    """Cover guesses and sheet count for the upload form. Unlike a full read
    it never fails for want of items, so a layout the rules don't know can
    still be uploaded and read with AI help."""
    sheets = visible_sheets(open_workbook(path))
    cover_sheet = next((ws for ws in sheets if find_header(ws) is None), None)
    return (_cover_from(cover_sheet) if cover_sheet else Cover()), len(sheets)


def read_excel(path: str, column_mapping: dict | None = None) -> ExcelRead:
    wb = open_workbook(path)
    events: list[Event] = []
    issues: list[Issue] = []
    cover = Cover()
    cover_done = False
    sheets = visible_sheets(wb)

    last_cols: _Cols | None = None
    for sheet_no, ws in enumerate(sheets, start=1):
        cols = find_header(ws, mapping=column_mapping)
        if last_cols is not None and column_mapping is None:
            # No real header of its own (at best a row of working columns),
            # but laid out like the sheet before it.
            borrowed = replace(last_cols, header_row=0)
            fit = _fit_count(ws, borrowed)
            if fit >= 3 and (cols is None or fit > 2 * _fit_count(ws, cols)):
                cols = borrowed
        if cols is not None and cols.header_row == 0:
            issues.append(Issue(code="LAYOUT_BORROWED", severity="info", message=f'Sheet "{ws.title.strip()}" has no column headers. It was read with the columns of the sheet before it — check its items.'))
        if cols is None:
            if not cover_done:
                cover = _cover_from(ws)
                cover_done = True
            if _item_like_rows(ws) >= 20:
                issues.append(Issue(code="SHEET_SKIPPED", severity="warning", message=f'Sheet "{ws.title.strip()}" looks like it has BOQ items but no column headers, so it wasn\'t read.'))
            continue
        last_cols = cols
        header_desc = _text(ws.cell(cols.header_row, cols.desc).value).upper() if cols.header_row else ""

        bill_no, title = str(sheet_no).zfill(2), ws.title.strip()
        for row in ws.iter_rows(min_row=1, max_row=cols.header_row - 1):
            for c in row:
                m = _BILL.search(_text(c.value))
                if m:
                    bill_no, title = m.group(1).zfill(2), m.group(2).strip()
        events.append(BillStart(bill_no=bill_no, title=title))

        block: list[tuple[str, bool]] = []
        pending_text: list[tuple[str, bool]] = []  # desc-only rows right above a possible item

        def flush() -> None:
            nonlocal block
            if block:
                events.append(TextBlock(block, sheet_no))
                block = []

        for row in ws.iter_rows(min_row=cols.header_row + 1):
            cell = {c.column: c for c in row}

            def val(col: int | None):
                return cell[col].value if col and col in cell else None

            desc_cell = cell.get(cols.desc)
            desc = _text(val(cols.desc))
            ref = _text(val(cols.ref)) if cols.ref else ""
            if ref == "0":
                ref = ""
            qty, unit = _num(val(cols.qty)), _text(val(cols.unit))
            rate, amount = _num(val(cols.rate)), _num(val(cols.amount))
            bold = bool(desc_cell is not None and desc_cell.font is not None and desc_cell.font.b)

            if not (desc or ref or qty or unit or rate or amount):
                block.extend(pending_text)
                pending_text = []
                flush()
                continue
            words = [w for w in (_text(c.value) for c in row) if w]
            if any(_PAGE_TOTAL.search(w) for w in words):
                events.append(TotalRow(amount_raw=amount, page=sheet_no))
                continue
            if any(_SKIP_ANY_CELL.search(w) for w in words) or (desc and _SKIP_ROW.search(desc) and not ref):
                continue
            if sum(1 for c in row if _text(c.value) == get_column_letter(c.column)) >= 3:
                continue  # a row of column letters (A, B, C…) above the table
            no_numbers = qty is None and rate is None and amount is None
            if no_numbers and (len(ref) > 12 or _WORD.search(ref)):
                desc, ref = desc or ref, ""  # a title typed in the reference column
            if ref and no_numbers and not (desc or unit):
                continue  # a lone reference letter with nothing else on the row
            if _WORD.search(ref) and not desc and qty is None:
                continue  # a summary line: a label in the reference column and a figure
            if header_desc and desc.upper() == header_desc:
                continue  # the column headers, repeated on every printed page
            if ref or qty is not None:
                lines = desc.splitlines() if desc else []
                if not lines and pending_text:
                    lines = [t for t, _ in pending_text]
                else:
                    block.extend(pending_text)
                pending_text = []
                flush()
                events.append(
                    ItemRow(ref=ref or "-", lines=lines, page=sheet_no, qty_raw=qty, unit_raw=unit or None, rate_raw=rate, amount_raw=amount)
                )
                continue
            # Description only.
            block.extend(pending_text)
            pending_text = [(line, bold) for line in desc.splitlines() if line.strip()]
        block.extend(pending_text)
        flush()

    if not any(isinstance(e, ItemRow) for e in events):
        raise ExtractError(
            "NO_ITEMS",
            "No BOQ items were found. Each bill sheet needs a header row with Description, Qty and Unit columns.",
        )
    return ExcelRead(events=events, cover=cover, page_totals=[], issues=issues, pages=len(sheets))
