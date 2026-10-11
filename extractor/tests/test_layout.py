"""AI layout reading (app/layout.py) and the Excel layouts it grew out of.

No AI is called here: tests label the rows themselves, exactly as the web
app would after asking a provider.
"""

import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import Font

from app.extract import extract
from app.hierarchy import BillStart, ItemRow, TotalRow
from app.layout import TOP, _cell_text, build_table, describe, mask
from app.main import app
from app.numbers import parse_number
from app.pdf_reader import ExtractError, Word
from app.validate import check_run_totals

from .synthetic_boq import Page, Row, draw

TOKEN = {"x-extractor-token": "secret-for-tests"}


def _items(result):
    return [(s, m, i) for b in result.bills for s in b.sections for m in s.main_descriptions for i in m.items]


def _leaks(described: dict) -> list[str]:
    """Any standalone number left in what would be sent to the AI."""
    return [
        tok
        for r in [*described["rows"], *({"cells": c} for l in described["layouts"] for c in l["sample"])]
        for v in r["cells"].values()
        for tok in v.split()
        if parse_number(tok.strip("()[]{},;:$£€").rstrip(".")) is not None
    ]


# ── Masking ───────────────────────────────────────────────────────────────


def test_mask_hides_standalone_numbers_only():
    assert mask("To balcony 4 m2 50.00 1,250.00") == "To balcony # m2 # #"
    assert mask("(120) 3,500.00; total: 10.") == "# # total: #"
    assert mask("500mm dia P4 piles, ref 2.5.1, 20000mm long") == "500mm dia P4 piles, ref 2.5.1, 20000mm long"
    assert mask("Rate Only") == "Rate Only"


def test_split_number_fragments_are_glued():
    def w(text, x0, x1):
        return Word(text, x0, x1, 0, 8, False)

    assert _cell_text([w("2", 548.1, 553.2), w("00.00", 553.1, 575.6)]) == "200.00"
    assert _cell_text([w("3", 540.0, 545.0), w(",500.00", 545.5, 570.0)]) == "3,500.00"
    assert _cell_text([w("Assumed", 600, 640), w("wood", 643, 665)]) == "Assumed wood"
    assert _cell_text([w("12", 500, 510), w("4", 520, 525)]) == "12 4"  # a real gap: two numbers


# ── Excel layouts the rules now read ──────────────────────────────────────


def _aldar_like(path: Path, second_sheet_header: bool = False) -> None:
    """Title block, a row of working columns, the real header far down,
    page footers in other columns, a repeated header and 0 refs."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Villa BoQ"
    ws.append(["A", "B", "C", "D", "E", "F", None, "ADJ. QTY", "UNIT", "Changes in Qty / Description"])
    ws["A2"], ws["A3"] = "CLIENT PJSC", "MARINA VILLAS"
    ws["A2"].font = Font(bold=True)
    for col, label in zip("ABCDEF", ["Ref", "Description", "QTY", "Unit", "Rate", "Total"]):
        ws[f"{col}55"] = label
    B = "bold"
    rows = [
        (None, "Section B - Site Works", None, None, None, None, B),
        None,
        ("0", "B9 Excavation", None, None, None, None, B),
        None,
        (None, "Excavation including all excavating, earthwork support and all associated works", None, None, None, None, None),
        None,
        ("A", "To receive pile caps", 100, "m3", 18, 1800, None),
        ("B", "To receive ground beams", 50, "m3", 18, 900, None),
        None,
        ("Section B/Page 1 of 2", None, None, "To Collection:", None, 2700, B),
        ("Ref", "Description", "QTY", "Unit", "Rate", "Total", B),
        (None, "B11 Disposal", None, None, None, None, B),
        None,
        ("C", "Disposal off site", 150, "m3", 10, 1500, None),
        ("Provisional Sums", None, None, None, None, 48875, None),
        ("Section B/Page 2 of 2", None, None, "To Collection:", None, 1500, B),
    ]
    for r, values in enumerate(rows, start=56):
        if values is None:
            continue  # a blank row between blocks
        for col, v in zip("ABCDEF", values):
            ws[f"{col}{r}"] = v
        if values[6]:
            ws[f"B{r}"].font = Font(bold=True)
    ws2 = wb.create_sheet("Boundary Wall")
    ws2.append(["A", "B", "C", "D", "E", "F"])
    ws2["A3"] = "Bills of Quantities (Boundary wall)"
    if second_sheet_header:
        ws2.append(["Ref", "Description", "QTY", "Unit", "Rate", "Total"])
    for values in [
        (None, "Front Wall", None, None, None, None),
        ("A", "Anti-termite treatment", 1805.41, "m2", 4, 7221.64),
        ("B", "Excavation for wall footing", 120, "m3", 18, 2160),
        ("C", "Blinding concrete", 30, "m3", 300, 9000),
    ]:
        ws2.append(list(values))
    wb.save(path)


def test_reads_header_far_down_and_skips_footers_and_titles(tmp_path):
    path = tmp_path / "aldar.xlsx"
    _aldar_like(path)
    result = extract(str(path), "xlsx")
    first = [i for s, m, i in _items(result) if i.page == 1]
    assert [(i.ref, i.description, i.qty, i.amount) for i in first] == [
        ("A", "To receive pile caps", "100", "1800"),
        ("B", "To receive ground beams", "50", "900"),
        ("C", "Disposal off site", "150", "1500"),
    ]
    s, m, _ = _items(result)[0]
    assert (s.parent_heading, s.heading) == ("Section B - Site Works", "B9 Excavation")
    assert m.text == "Excavation including all excavating, earthwork support and all associated works"
    assert not any(i.code == "TOTAL_MISMATCH" for i in result.issues)  # both page footers reconcile
    assert not result.quality.recommend_ai


def test_sheet_without_headers_borrows_the_previous_layout(tmp_path):
    path = tmp_path / "aldar.xlsx"
    _aldar_like(path)
    result = extract(str(path), "xlsx")
    assert [b.title for b in result.bills] == ["Villa BoQ", "Boundary Wall"]
    second = [i for s, m, i in _items(result) if i.page == 2]
    assert [(i.ref, i.qty, i.rate) for i in second] == [("A", "1805.41", "4"), ("B", "120", "18"), ("C", "30", "300")]
    assert any(i.code == "LAYOUT_BORROWED" for i in result.issues)


def test_page_total_mismatch_is_reported(tmp_path):
    events = [BillStart("01", "X"), ItemRow("A", ["a"], 1, amount_raw="100"), ItemRow("B", ["b"], 1, amount_raw="50"), TotalRow("200", 1), TotalRow("75", 2)]
    issues = check_run_totals(events)
    assert [i.code for i in issues] == ["TOTAL_MISMATCH"]  # the second total has no items before it
    assert "150.00" in issues[0].message and "200.00" in issues[0].message


def test_inspect_does_not_need_items(tmp_path):
    path = tmp_path / "odd.xlsx"
    wb = Workbook()
    wb.active.append(["Code", "Work", "Count"])
    wb.active.append(["A", "Pavers", 10])
    wb.save(path)
    client = TestClient(app)
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("EXTRACTOR_TOKEN", "secret-for-tests")
        r = client.post("/v1/inspect", headers=TOKEN, files={"file": ("odd.xlsx", path.read_bytes())})
    assert r.status_code == 200, r.text
    assert r.json()["pages"] == 1


# ── The masked row table ──────────────────────────────────────────────────


def _headerless(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Works"
    for values in [
        ("TOWER PROJECT", None, None, None, None, None),
        (None, "Paving", None, None, None, None),
        (None, "Interlocking pavers, 80mm thick, laid on sand bed", None, None, None, None),
        ("1", "To footpaths", 120, "m2", 85, 10200),
        ("2", "To driveways", 40, "m2", 95, 3800),
        (None, None, None, "Page total", None, 14000),
    ]:
        ws.append(list(values))
    ws["B2"].font = Font(bold=True)
    wb.save(path)


def test_row_table_masks_every_number(tmp_path):
    path = tmp_path / "plain.xlsx"
    _headerless(path)
    described = describe(build_table(str(path), "xlsx"))
    assert described["layouts"][0]["roles"] is None  # no header: the AI is asked
    assert described["layouts"][0]["columns"] == ["A", "B", "C", "D", "E", "F"]
    row = described["rows"][3]["cells"]
    assert row == {"A": "#", "B": "To footpaths", "C": "#", "D": "m2", "E": "#", "F": "#"}
    assert not _leaks(described)
    assert len(described["rows_hash"]) == 64


def _plan(described: dict, roles: dict, labels: str) -> dict:
    return {"rows_hash": described["rows_hash"], "roles": {l["key"]: roles for l in described["layouts"]}, "labels": labels}


EXCEL_ROLES = {"reference": "A", "description": "B", "quantity": "C", "unit": "D", "rate": "E", "amount": "F"}


def test_rebuilds_boq_from_labels_with_numbers_from_the_file(tmp_path):
    path = tmp_path / "plain.xlsx"
    _headerless(path)
    described = describe(build_table(str(path), "xlsx"))
    result = extract(str(path), "xlsx", layout=_plan(described, EXCEL_ROLES, "XHMIIT"))
    assert result.parser == "ai-layout"
    (s, m, a), (_, _, b) = _items(result)
    assert (s.heading, m.text) == ("Paving", "Interlocking pavers, 80mm thick, laid on sand bed")
    assert (a.ref, a.description, a.qty, a.unit, a.rate, a.amount) == ("1", "To footpaths", "120", "m²", "85", "10200")
    assert (b.ref, b.amount) == ("2", "3800")
    assert result.issues[0].code == "AI_LAYOUT"
    assert not any(i.code == "TOTAL_MISMATCH" for i in result.issues)


def test_short_main_description_stays_a_main_description(tmp_path):
    path = tmp_path / "plain.xlsx"
    _headerless(path)
    described = describe(build_table(str(path), "xlsx"))
    # Row 1 ("TOWER PROJECT") labelled M: short and bold-free, but the AI said
    # it's a main description, so it must not be turned into a heading.
    result = extract(str(path), "xlsx", layout=_plan(described, EXCEL_ROLES, "MXXIIT"))
    s, m, _ = _items(result)[0]
    assert s.heading is None and m.text == "TOWER PROJECT"


def test_continuation_lines_and_mislabelled_priced_rows(tmp_path):
    path = tmp_path / "plain.xlsx"
    _headerless(path)
    described = describe(build_table(str(path), "xlsx"))
    # "Interlocking pavers…" as the item's first line, "To footpaths" (with
    # the numbers) as its continuation; the second priced row wrongly ignored.
    result = extract(str(path), "xlsx", layout=_plan(described, EXCEL_ROLES, "XHICXT"))
    ((_, _, item),) = _items(result)
    assert item.description == "Interlocking pavers, 80mm thick, laid on sand bed To footpaths"
    assert (item.ref, item.qty, item.amount) == ("1", "120", "10200")
    codes = [i.code for i in result.issues]
    assert "PRICED_ROW_SKIPPED" in codes and "TOTAL_MISMATCH" in codes


@pytest.mark.parametrize(
    "change, message",
    [
        (lambda p: p.update(rows_hash="0" * 64), "read differently"),
        (lambda p: p.update(labels=p["labels"][:-1]), "labels"),
        (lambda p: p.update(labels=p["labels"][:-1] + "Z"), "labels"),
        (lambda p: [r.update(quantity="Q") for r in p["roles"].values()], "column roles"),
        (lambda p: [r.update(unit="B") for r in p["roles"].values()], "column roles"),  # same column twice
        (lambda p: [r.update(description=None) for r in p["roles"].values()], "description"),
    ],
)
def test_rejects_plans_that_dont_fit_the_file(tmp_path, change, message):
    path = tmp_path / "plain.xlsx"
    _headerless(path)
    described = describe(build_table(str(path), "xlsx"))
    plan = json.loads(json.dumps(_plan(described, EXCEL_ROLES, "XHMIIT")))
    change(plan)
    with pytest.raises(ExtractError) as err:
        extract(str(path), "xlsx", layout=plan)
    assert err.value.code == "INVALID_LAYOUT" and message in err.value.message


def _pdf(path: Path) -> None:
    draw(
        str(path),
        pages=[
            Page(
                bill="BILL NO. 02 - HARD LANDSCAPE",
                rows=[
                    Row("heading", ["Paving"]),
                    Row("text", ["Supply and lay concrete pavers"]),
                    Row("item", ["60mm thick"], ref="A", qty="1,200", unit="m2", rate="85.00", amount="102,000.00"),
                    Row("item", ["80mm thick"], ref="B", qty="300", unit="m2", rate="95.50", amount="28,650.00"),
                ],
                footer_total="130,650.00",
            )
        ],
    )


def test_pdf_row_table_uses_header_columns(tmp_path):
    path = tmp_path / "b.pdf"
    _pdf(path)
    table = build_table(str(path), "pdf")
    described = describe(table)
    (layout,) = described["layouts"]
    assert layout["roles"] == {"reference": "c0", "description": "c1", "quantity": "c2", "unit": "c3", "rate": "c4", "amount": "c5"}
    assert not _leaks(described)
    item_a = next(r for r in table.rows if r.cells.get("c0") == "A")
    assert (item_a.cells["c2"], item_a.cells["c3"], item_a.cells["c4"], item_a.cells["c5"]) == ("1,200", "m2", "85.00", "102,000.00")
    assert any(TOP in r.cells and "HARD LANDSCAPE" in r.cells[TOP] for r in table.rows)


def test_pdf_rebuilt_from_labels(tmp_path):
    path = tmp_path / "b.pdf"
    _pdf(path)
    described = describe(build_table(str(path), "pdf"))
    labels = []
    for r in described["rows"]:
        cells = r["cells"]
        text = " ".join(cells.values())
        if TOP in cells:
            labels.append("B" if "BILL" in text else "X")
        elif "COLLECTION" in text.upper() or list(cells) == ["c5"]:
            labels.append("T")  # the footer's label and its figure sit on separate lines
        elif cells.get("c0"):
            labels.append("I")
        elif "Paving" in text:
            labels.append("H")
        elif "pavers" in text:
            labels.append("M")
        else:
            labels.append("X")
    roles = described["layouts"][0]["roles"]
    result = extract(str(path), "pdf", layout={"rows_hash": described["rows_hash"], "roles": {described["layouts"][0]["key"]: roles}, "labels": "".join(labels)})
    assert [(b.bill_no, b.title) for b in result.bills] == [("02", "HARD LANDSCAPE")]
    (s, m, a), (_, _, b) = _items(result)
    assert (s.heading, m.text) == ("Paving", "Supply and lay concrete pavers")
    assert (a.ref, a.qty, a.unit, a.rate, a.amount) == ("A", "1200", "m²", "85", "102000")
    assert (b.ref, b.qty, b.rate, b.amount) == ("B", "300", "95.5", "28650")
    assert not any(i.code == "TOTAL_MISMATCH" for i in result.issues)
    # The totals check is live: drop item B and the footer no longer adds up.
    broken = "".join("X" if r["cells"].get("c0") == "B" else lab for r, lab in zip(described["rows"], labels))
    result = extract(str(path), "pdf", layout={"rows_hash": described["rows_hash"], "roles": {described["layouts"][0]["key"]: roles}, "labels": broken})
    assert [i.code for i in result.issues if i.code != "AI_LAYOUT"] == ["PRICED_ROW_SKIPPED", "TOTAL_MISMATCH"]


def test_rows_and_layout_endpoints(tmp_path, monkeypatch):
    monkeypatch.setenv("EXTRACTOR_TOKEN", "secret-for-tests")
    client = TestClient(app)
    path = tmp_path / "plain.xlsx"
    _headerless(path)
    r = client.post("/v1/rows", headers=TOKEN, files={"file": ("plain.xlsx", path.read_bytes())})
    assert r.status_code == 200, r.text
    described = r.json()
    assert not _leaks(described)
    plan = _plan(described, EXCEL_ROLES, "XHMIIT")
    r = client.post("/v1/extract", headers=TOKEN, files={"file": ("plain.xlsx", path.read_bytes())}, data={"layout": json.dumps(plan)})
    assert r.status_code == 200, r.text
    assert r.json()["parser"] == "ai-layout" and r.json()["stats"]["items"] == 2
    r = client.post("/v1/extract", headers=TOKEN, files={"file": ("plain.xlsx", path.read_bytes())}, data={"layout": "{not json"})
    assert r.status_code == 422 and r.json()["code"] == "INVALID_LAYOUT"


def test_weak_rule_based_reading_recommends_ai(tmp_path):
    path = tmp_path / "plain.xlsx"
    _headerless(path)
    with pytest.raises(ExtractError) as err:
        extract(str(path), "xlsx")  # no header row anywhere: the rules find nothing
    assert err.value.code == "NO_ITEMS"
