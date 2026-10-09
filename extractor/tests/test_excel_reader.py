from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from app.extract import extract
from app.pdf_reader import ExtractError


def make_boq(path: Path) -> None:
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["B5"] = "MARINA TOWER"
    cover["B7"] = "Design Development 50%"
    cover["B8"] = "Date: 12/03/2026"

    ws = wb.create_sheet("Concrete")
    ws["A1"] = "BILL NO. 03 - CONCRETE WORKS"
    for col, label in zip("ABCDEF", ["ITEM", "DESCRIPTION", "QTY", "UNIT", "RATE", "AMOUNT"]):
        ws[f"{col}3"] = label
    rows = [
        (None, "In-situ concrete", None, None, None, None, True),
        (None, "Grade C40 concrete, 20mm aggregate, in:", None, None, None, None, False),
        ("A", "Column bases", 45, "m3", 850, 38250, False),
        ("B", "Ground beams", 30, "m3", 875.5, 26265, False),
        (None, None, None, None, None, None, False),
        (None, "Formwork", None, None, None, None, True),
        ("C", "Sides of ground beams", 120, "m2", 65, 7800, False),
        (None, "CARRIED TO COLLECTION", None, None, None, 72315, False),
    ]
    for i, (ref, desc, qty, unit, rate, amount, bold) in enumerate(rows, start=4):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"], ws[f"E{i}"], ws[f"F{i}"] = ref, desc, qty, unit, rate, amount
        if bold:
            ws[f"B{i}"].font = Font(bold=True)
    wb.save(path)


def test_reads_excel_boq(tmp_path: Path):
    path = tmp_path / "boq.xlsx"
    make_boq(path)
    result = extract(str(path), "xlsx")

    assert result.file_type == "xlsx"
    assert [b.bill_no for b in result.bills] == ["03"]
    assert result.bills[0].title == "CONCRETE WORKS"
    items = [(s, m, i) for b in result.bills for s in b.sections for m in s.main_descriptions for i in m.items]
    assert [i.ref for *_, i in items] == ["A", "B", "C"]

    s, m, a = items[0]
    assert s.heading == "In-situ concrete"
    assert m.text == "Grade C40 concrete, 20mm aggregate, in:"
    assert (a.qty, a.unit, a.rate, a.amount) == ("45", "m³", "850", "38250")
    assert not a.flags
    assert items[1][2].rate == "875.5"

    s, m, c = items[2]
    assert s.heading == "Formwork" and m.text == ""
    assert result.stats.priced_items == 3


def test_excel_cover_guesses(tmp_path: Path):
    path = tmp_path / "boq.xlsx"
    make_boq(path)
    cover = extract(str(path), "xlsx").cover
    assert cover.project_name == "MARINA TOWER"
    assert cover.stage_guess == "DD 50%"
    assert cover.boq_date == "2026-03-12"


def test_excel_without_items_is_rejected(tmp_path: Path):
    path = tmp_path / "empty.xlsx"
    Workbook().save(path)
    with pytest.raises(ExtractError) as err:
        extract(str(path), "xlsx")
    assert err.value.code == "NO_ITEMS"
